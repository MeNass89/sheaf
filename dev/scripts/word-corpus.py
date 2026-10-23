#!/usr/bin/env python3
"""PDF to Word corpus (F-005): Word originals made here, turned into PDFs by LibreOffice, exported back
to Word by Sheaf, and each export measured against its original.

    python3 dev/scripts/word-corpus.py [path/to/sheaf]     # default: the debug CLI

Needs LibreOffice (soffice) and python-docx (installed on first run in ~/.cache/sheaf-corpus-venv).
Writes into dev/.tmp/word-corpus (hygiene deletes it). Prints one line per document and check, then a
summary; exits 1 if a check fails.
"""
import os, re, subprocess, sys, zipfile

VENV = os.path.expanduser("~/.cache/sheaf-corpus-venv")
if sys.prefix != VENV:
    if not os.path.exists(f"{VENV}/bin/python"):
        subprocess.run([sys.executable, "-m", "venv", VENV], check=True)
        subprocess.run([f"{VENV}/bin/pip", "-q", "install", "python-docx"], check=True)
    os.execv(f"{VENV}/bin/python", [f"{VENV}/bin/python", *sys.argv])

import docx  # noqa: E402
from docx.enum.text import WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt, RGBColor, Inches  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(ROOT, "dev/.tmp/word-corpus")
SHEAF = sys.argv[1] if len(sys.argv) > 1 else os.path.expanduser("~/.cache/cargo-target/sheaf/debug/sheaf")
LOREM = ("Sheaf keeps every document on this computer. Nothing is uploaded, nothing is tracked, and the "
         "files stay where you put them. ")


def fill(cell, hex_colour):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_colour)
    cell._tc.get_or_add_tcPr().append(shd)


def letter(path):
    d = docx.Document()
    d.add_paragraph("Dr Jane Doe\n12 Rue des Lilas\n75011 Paris")
    d.add_paragraph("Paris, 23 October 2026")
    p = d.add_paragraph("Dear colleague, ")
    p.add_run("thank you").bold = True
    p.add_run(" for your referral. " + LOREM * 3)
    d.add_paragraph(LOREM * 4)
    d.add_paragraph("Yours sincerely,")
    d.save(path)


def fiche(path):
    d = docx.Document()
    d.add_heading("ITEM 282: Crohn disease", 1)
    t = d.add_table(rows=4, cols=3)
    t.style = "Table Grid"
    t.cell(0, 0).merge(t.cell(0, 2)).text = "MALADIE DE CROHN"
    fill(t.cell(0, 0), "77933C")
    t.cell(1, 0).merge(t.cell(2, 0)).text = "Clinique"
    fill(t.cell(1, 0), "D7E4BD")
    t.cell(1, 1).text = "Symptoms"
    # A word in another style mid-line: the space before it must stay (a lost one merges words).
    cp = t.cell(1, 2).paragraphs[0]
    cp.add_run("Abdominal pain, ")
    cp.add_run("diarrhoea").bold = True
    cp.add_run(", weight loss")
    t.cell(2, 1).text = "Signs"
    t.cell(2, 2).text = "Fever, growth delay"
    t.cell(3, 0).text = "Bio"
    t.cell(3, 1).text = "CRP"
    t.cell(3, 2).text = "Raised in flares"
    p = d.add_paragraph("Evolution by ")
    r = p.add_run("flares")
    r.bold = True
    r.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    p.add_run(" separated by remissions.")
    d.save(path)


def invoice(path):
    d = docx.Document()
    d.add_heading("Invoice 2026-114", 1)
    d.add_paragraph("Billed to: Clinique du Parc")
    t = d.add_table(rows=1, cols=4)
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ["Item", "Qty", "Price", "Total"]):
        c.text = h
    for item, q, pr in [("Consultation", 2, 50), ("Ultrasound", 1, 120), ("Report", 1, 30)]:
        r = t.add_row().cells
        r[0].text, r[1].text, r[2].text, r[3].text = item, str(q), f"{pr}.00", f"{q * pr}.00"
    d.add_paragraph("Total due: 250.00 EUR")
    d.save(path)


def columns(path):
    d = docx.Document()
    d.add_heading("Two columns", 1)
    sect = d.sections[0]._sectPr
    cols = OxmlElement("w:cols")
    cols.set(qn("w:num"), "2")
    sect.append(cols)
    d.add_paragraph("First column. " + LOREM * 6)
    d.add_paragraph("Second part. " + LOREM * 6)
    d.save(path)


def report(path):
    d = docx.Document()
    d.add_heading("Annual report", 1)
    d.add_heading("Summary", 2)
    p = d.add_paragraph(LOREM * 2)
    r = p.add_run("Red words.")
    r.font.color.rgb = RGBColor(0xC0, 0, 0)
    r.italic = True
    d.add_heading("Details", 2)
    d.add_paragraph(LOREM * 5)
    d.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)
    d.add_heading("Appendix", 2)
    d.add_paragraph(LOREM)
    d.save(path)


DOCS = {"letter": letter, "fiche": fiche, "invoice": invoice, "columns": columns, "report": report}


def soffice(*args, outdir):
    subprocess.run(["soffice", "--headless", "--norestore", *args, "--outdir", outdir],
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180, env={**os.environ, "HOME": OUT})


def text_of(el):
    """An element's text, a line break or a tab being a space."""
    return "".join((t.text or "") if t.tag == qn("w:t") else " " for t in el.iter(qn("w:t"), qn("w:br"), qn("w:tab")))


def words(path):
    """The document's words in order: paragraphs, then table cells row by row, where they are."""
    d = docx.Document(path)
    out = []
    body = d.element.body
    for el in body.iterchildren():
        if el.tag == qn("w:p"):
            out.append(text_of(el))
        elif el.tag == qn("w:tbl"):
            seen = set()
            for tc in el.iter(qn("w:tc")):
                if id(tc) in seen:
                    continue
                seen.add(id(tc))
                out.append(text_of(tc))
    return re.findall(r"\w+", " ".join(out).lower())


def tables(path):
    d = docx.Document(path)
    return [(len(t.rows), sum(len({id(c._tc) for c in r.cells}) for r in t.rows)) for t in d.tables]


def pages(pdf):
    out = subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout
    m = re.search(r"Pages:\s+(\d+)", out)
    return int(m.group(1)) if m else 0


def main():
    os.makedirs(OUT, exist_ok=True)
    fails = 0
    for name, make in DOCS.items():
        orig = f"{OUT}/{name}.docx"
        make(orig)
        soffice("--convert-to", "pdf", orig, outdir=OUT)
        pdf = f"{OUT}/{name}.pdf"
        export = f"{OUT}/{name}-sheaf.docx"
        if os.path.exists(export):
            os.remove(export)
        r = subprocess.run([SHEAF, "to-word", pdf, "-o", export], capture_output=True, text=True)
        if r.returncode != 0 or not os.path.exists(export):
            print(f"FAIL {name}: export failed: {r.stderr.strip()[:200]}")
            fails += 1
            continue
        back = f"{OUT}/back"
        os.makedirs(back, exist_ok=True)
        soffice("--convert-to", "pdf", export, outdir=back)
        xml = zipfile.ZipFile(export).read("word/document.xml").decode("utf8", "replace")
        pua = [f"U+{ord(c):04X}" for c in xml if 0xE000 <= ord(c) <= 0xF8FF]
        boxes = len(re.findall(r"<wp:anchor|<w:framePr|<v:textbox|<wps:txbx", xml))
        w0, w1 = words(orig), words(export)
        checks = [
            ("words in order", w0 == w1, f"{len(w1)} of {len(w0)} words" + ("" if w0 == w1 else f", first difference at {next((i for i, (a, b) in enumerate(zip(w0, w1)) if a != b), min(len(w0), len(w1)))}")),
            ("tables", tables(orig) == tables(export), f"{tables(export)} for {tables(orig)} (rows, cells)"),
            ("no floating text boxes", boxes == 0, f"{boxes}"),
            ("no private-use characters", not pua, f"{len(pua)}" + (f": {sorted(set(pua))[:5]}" if pua else "")),
            ("page count", abs(pages(f'{back}/{name}-sheaf.pdf') - pages(pdf)) <= 1, f"{pages(f'{back}/{name}-sheaf.pdf')} for {pages(pdf)}"),
        ]
        for what, ok, detail in checks:
            print(f"{'PASS' if ok else 'FAIL'} {name}: {what} ({detail})")
            fails += not ok
    print(f"--- {fails} check(s) failed" if fails else "--- all checks pass")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
